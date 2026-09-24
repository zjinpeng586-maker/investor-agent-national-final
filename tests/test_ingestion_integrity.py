"""User-reported import failures: real PDFs, units, ownership, and provenance."""
from __future__ import annotations

from copy import deepcopy
import io
import json
from hashlib import sha256
from pathlib import Path

import pytest
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.platypus import Table, TableStyle

from core.parsers import parse_pdf, parse_tabular, _str_from_context
from core.service import prepare_import, commit_import, ingest_pdf_file, save_upload
from core.storage import workspace_context
from core import db


def make_pdf(*, table: bool = False, units: bool = True, duplicate: bool = False,
             comparative: bool = False, half_year: bool = False) -> bytes:
    pdfmetrics.registerFont(UnicodeCIDFont('STSong-Light'))
    buffer = io.BytesIO()
    doc = canvas.Canvas(buffer, pagesize=(800, 700))
    doc.setFont('STSong-Light', 12)
    doc.drawString(30, 670, '测试股份有限公司')
    doc.drawString(30, 645, '2024年半年度报告' if half_year else '2024年年度报告')
    doc.drawString(30, 620, '主要会计数据和财务指标')
    if units:
        doc.drawString(30, 595, '单位：亿元')
    rows = [['指标', '2024年', '2023年']] if comparative else [['指标', '2024年']]
    rows += [['营业收入', '100', '90'], ['归母净利润', '10', '8'], ['经营现金流', '20', '15']]
    if not comparative:
        rows = [row[:2] for row in rows]
    if table:
        financial = Table(rows, colWidths=[250] + [150] * (len(rows[0]) - 1), rowHeights=35)
        financial.setStyle(TableStyle([('FONTNAME', (0, 0), (-1, -1), 'STSong-Light'),
                                      ('FONTSIZE', (0, 0), (-1, -1), 12),
                                      ('GRID', (0, 0), (-1, -1), 1, 'black')]))
        financial.wrapOn(doc, 700, 400)
        financial.drawOn(doc, 30, 420)
    else:
        y = 565
        for row in rows:
            doc.drawString(30, y, '      '.join(row))
            y -= 30
    doc.drawString(30, 360, '审计意见：无保留意见')
    if duplicate:
        doc.drawString(30, 325, '营业收入 200')
    doc.save()
    return buffer.getvalue()


@pytest.fixture
def isolated(tmp_path):
    with workspace_context('evaluation', tmp_path, allow_writes=True):
        db.init_db()
        yield tmp_path


@pytest.mark.parametrize('table', [False, True])
def test_pdf_rows_are_not_previous_metric_values(table):
    company, records, warnings, _ = parse_pdf('测试2024年年度报告.pdf', make_pdf(table=table))
    assert company == '测试股份有限公司'
    assert len(records) == 1
    record = records[0]
    assert (record['revenue'], record['net_profit'], record['operating_cashflow']) == (100, 10, 20)
    assert not record['_errors'], warnings
    assert record['audit_opinion'] == '无保留意见'
    assert record['_provenance']['revenue']['page'] == 1
    assert record['_provenance']['net_profit']['raw_value'] == '10'
    assert record['_provenance']['operating_cashflow']['cell'] != record['_provenance']['revenue']['cell']


@pytest.mark.parametrize('opinion', ['无保留意见', '标准的无保留意见', '带强调事项段的无保留意见', '保留意见', '否定意见', '无法表示意见'])
def test_audit_negation_is_preserved(opinion):
    assert _str_from_context('审计意见：' + opinion, ['审计意见']) == opinion


@pytest.mark.parametrize('table', [False, True])
def test_pdf_year_columns_match_exactly(table):
    _, records, warnings, _ = parse_pdf('2024年年度报告.pdf', make_pdf(table=table, comparative=True))
    rows = {record['year']: record for record in records}
    assert rows[2024]['revenue'] == 100
    assert rows[2023]['revenue'] == 90
    assert rows[2023]['net_profit'] == 8
    assert rows[2023]['operating_cashflow'] == 15
    assert not any(record['_errors'] for record in records), warnings


def test_pdf_unknown_units_block_even_confirmed(isolated):
    batch = prepare_import('2024年年度报告.pdf', make_pdf(units=False))
    assert not batch['can_commit']
    assert any('单位未知' in error for error in batch['errors'])
    assert not commit_import(batch, confirmed=True)['ok']
    assert len(db.fetch_companies()) == 0
    assert not list(isolated.rglob('*.pdf'))


def test_pdf_manual_unit_resolves_unknown_unit():
    batch = prepare_import('2024年年度报告.pdf', make_pdf(units=False), money_unit='亿元')
    assert batch['can_commit'], batch['errors']


def test_pdf_ambiguous_rows_block():
    batch = prepare_import('2024年年度报告.pdf', make_pdf(duplicate=True))
    assert not batch['can_commit']
    assert any('数值冲突' in error for error in batch['errors'])


def test_year_conflict_and_half_year_do_not_relabel():
    batch = prepare_import('2024年年度报告.pdf', make_pdf(), year_hint=2025)
    assert not batch['can_commit']
    assert any('冲突' in error for error in batch['errors'])
    half_year = prepare_import('半年.pdf', make_pdf(half_year=True))
    assert not half_year['can_commit']
    assert any('半年/季度' in error for error in half_year['errors'])
    wrong_name = prepare_import('2025年年度报告.pdf', make_pdf())
    assert not wrong_name['can_commit']
    assert any('文件名年度' in error for error in wrong_name['errors'])


def test_csv_yuan_and_ratio_are_normalized_per_column():
    content = '企业名称,年份,营业收入（元）,归母净利润（万元）,经营现金流（亿元）,ROE（比例）,资产负债率（%）\n甲公司,2024,10000000000,100000,20,0.12,60\n'.encode()
    batch = prepare_import('单位.csv', content)
    assert batch['can_commit'], batch['errors']
    row = batch['records'][0]
    assert (row['revenue'], row['net_profit'], row['operating_cashflow'], row['roe'], row['debt_ratio']) == (100, 10, 20, 12, 60)
    assert row['_provenance']['revenue']['raw_unit'] == '元'
    assert row['_provenance']['revenue']['cell'] == 'CSV!C2'


def test_csv_multiple_companies_keep_ownership(isolated):
    content = '企业名称,年份,营业收入（亿元）\n甲公司,2024,100\n乙公司,2024,200\n'.encode()
    batch = prepare_import('两家公司.csv', content)
    assert batch['can_commit'], batch['errors']
    result = commit_import(batch, confirmed=True)
    assert result['ok']
    companies = {row['name']: row['id'] for row in db.fetch_companies()}
    assert set(companies) == {'甲公司', '乙公司'}
    assert db.fetch_company_metrics(companies['乙公司'])[0]['revenue'] == 200
    conflict = prepare_import('两家公司.csv', content, manual_company_name='甲公司')
    assert not conflict['can_commit']
    assert any('企业归属冲突' in error for error in conflict['errors'])


def test_unknown_percent_and_currency_do_not_guess():
    content = '公司名称,年份,营业收入,ROE\n甲公司,2024,100,0.12\n'.encode()
    batch = prepare_import('未知单位.csv', content)
    assert not batch['can_commit']
    fixed = prepare_import('未知单位.csv', content, money_unit='亿元', percent_unit='比例')
    assert fixed['can_commit'], fixed['errors']
    assert fixed['records'][0]['roe'] == 12
    foreign = prepare_import('美元.csv', '公司名称,年份,营业收入（美元）\n甲公司,2024,100\n'.encode(), money_unit='亿元')
    assert not foreign['can_commit']


def test_excel_percent_format_and_multiple_sheets():
    from openpyxl import Workbook
    book = Workbook()
    for sheet, name, revenue in [(book.active, '甲公司', 100), (book.create_sheet('第二家'), '乙公司', 200)]:
        sheet.append(['公司名称', '年份', '营业收入（亿元）', 'ROE'])
        sheet.append([name, 2024, revenue, .125])
        sheet['D2'].number_format = '0.0%'
    buffer = io.BytesIO()
    book.save(buffer)
    batch = prepare_import('两张表.xlsx', buffer.getvalue())
    assert batch['can_commit'], batch['errors']
    assert set(batch['companies']) == {'甲公司', '乙公司'}
    assert [row['roe'] for row in batch['records']] == [12.5, 12.5]


def test_excel_literal_percent_symbol_does_not_multiply_again():
    from openpyxl import Workbook
    book = Workbook()
    book.active.append(['公司名称', '年份', 'ROE（百分数）'])
    book.active.append(['甲公司', 2024, 12.5])
    book.active['C2'].number_format = '0.0"%"'
    buffer = io.BytesIO()
    book.save(buffer)
    batch = prepare_import('字面百分号.xlsx', buffer.getvalue())
    assert batch['can_commit'], batch['errors']
    assert batch['records'][0]['roe'] == 12.5


def test_duplicate_ownership_and_invalid_number_block():
    duplicated = prepare_import('重复.csv', '公司,年份,营业收入（亿元）\n甲,2024,100\n甲,2024,200'.encode())
    assert not duplicated['can_commit']
    invalid = prepare_import('错误.csv', '公司,年份,营业收入（亿元）\n甲,2024,约100至200'.encode())
    assert not invalid['can_commit']
    quarter = prepare_import('季度.csv', '公司,年份,营业收入（亿元）\n甲,2024上半年,100'.encode())
    assert not quarter['can_commit']


def test_preview_required_and_no_data_overwrite_by_default(isolated):
    old = prepare_import('原始.csv', '公司,年份,营业收入（亿元）,归母净利润（亿元）\n甲,2024,100,10'.encode())
    assert not commit_import(old)['ok']
    assert len(db.fetch_companies()) == 0
    assert commit_import(old, confirmed=True)['ok']
    incoming = prepare_import('更新.csv', '公司,年份,营业收入（亿元）,经营现金流（亿元）\n甲,2024,200,20'.encode())
    result = commit_import(incoming, confirmed=True)
    assert result['ok']
    cid = db.fetch_companies()[0]['id']
    row = db.fetch_company_metrics(cid)[0]
    assert (row['revenue'], row['net_profit'], row['operating_cashflow']) == (100, 10, 20)
    assert result['skipped_metrics'] == ['甲 2024 revenue']


def test_partial_update_preserves_other_metric_provenance(isolated):
    first = prepare_import('第一版.csv', '公司,年份,营业收入（亿元）,归母净利润（亿元）\n甲,2024,100,10'.encode())
    assert commit_import(first, confirmed=True)['ok']
    second = prepare_import('第二版.csv', '公司,年份,归母净利润（万元）\n甲,2024,110000'.encode())
    assert commit_import(second, confirmed=True, overwrite_existing=True)['ok']
    cid = db.fetch_companies()[0]['id']
    sources = {row['metric']: dict(row) for row in db.fetch_metric_sources(cid)}
    assert sources['revenue']['file_name'] == '第一版.csv'
    assert sources['net_profit']['file_name'] == '第二版.csv'
    assert sources['net_profit']['raw_unit'] == '万元'
    assert sources['revenue']['import_id'] != sources['net_profit']['import_id']


def test_no_builtin_fallback_and_legacy_requires_confirmation(isolated):
    with pytest.raises(ValueError, match='预览'):
        ingest_pdf_file('2024年年度报告.pdf', make_pdf())
    assert len(db.fetch_companies()) == 0
    missing = prepare_import('无单位.pdf', make_pdf(units=False), stock_code='002594', year_hint=2024)
    assert not missing['can_commit']
    assert not any(row.get('roe') for row in missing['records'])


def test_same_file_names_are_immutable_and_input_path_sanitized(isolated):
    first = save_upload('../../same.pdf', b'first')
    second = save_upload('..\\same.pdf', b'second')
    assert first.name == second.name == 'same.pdf'
    assert first != second
    assert first.read_bytes() == b'first'
    assert first.is_relative_to(isolated)


def test_tampered_preview_is_rejected(isolated):
    batch = prepare_import('原始.csv', '公司,年份,营业收入（亿元）\n甲,2024,100'.encode())
    bad = deepcopy(batch)
    bad['records'][0]['revenue'] = 999
    assert not commit_import(bad, confirmed=True)['ok']
    assert len(db.fetch_companies()) == 0


def test_missing_company_cell_is_not_inherited_from_first_row():
    batch = prepare_import('缺公司.csv', '公司,年份,营业收入（亿元）\n甲,2024,100\n,2023,90'.encode())
    assert not batch['can_commit']
    assert batch['records'][1]['company_name'] is None


def test_period_column_is_not_silently_annualized():
    batch = prepare_import('报告期.csv', '公司,年份,期间,营业收入（亿元）\n甲,2024,上半年,100'.encode())
    assert not batch['can_commit']
    assert any('期间' in error for error in batch['errors'])


def test_per_row_units_can_differ():
    batch = prepare_import('行单位.csv', '公司,年份,营业收入,金额单位\n甲,2024,100000000,元\n乙,2024,10,亿元'.encode())
    assert batch['can_commit'], batch['errors']
    assert [row['revenue'] for row in batch['records']] == [1, 10]


def test_money_never_accepts_a_percentage_cell():
    batch = prepare_import('错列.csv', '公司,年份,营业收入（元）\n甲,2024,29.02%'.encode())
    assert not batch['can_commit']
    assert any('疑似误取增减幅列' in error for error in batch['errors'])


def test_transaction_rolls_back_all_companies_on_failure(isolated, monkeypatch):
    from core import service
    batch = prepare_import('两家公司.csv', '公司,年份,营业收入（亿元）\n甲,2024,100\n乙,2024,200'.encode())
    actual = service.upsert_metric
    calls = []
    def fail_second(company_id, record, **kwargs):
        calls.append(company_id)
        if len(calls) == 2:
            raise RuntimeError('模拟第二行写入失败')
        actual(company_id, record, **kwargs)
    monkeypatch.setattr(service, 'upsert_metric', fail_second)
    result = commit_import(batch, confirmed=True)
    assert result['status'] == 'failed'
    assert not result['ok'] and not result['committed']
    assert '第二行' in result['error_message']
    assert not db.fetch_companies()
    assert not list(isolated.rglob('*.csv'))


def test_upload_size_limit_and_zip_expansion_are_blockers():
    batch = prepare_import('过大.csv', b'x' * (30 * 1024 * 1024 + 1))
    assert not batch['can_commit']
    assert any('30 MB' in error for error in batch['errors'])
    import zipfile
    compressed = io.BytesIO()
    with zipfile.ZipFile(compressed, 'w', compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr('xl/worksheets/sheet1.xml', b'x' * (2 * 1024 * 1024))
    bomb = prepare_import('压缩异常.xlsx', compressed.getvalue())
    assert not bomb['can_commit']
    assert any('压缩展开比例' in error for error in bomb['errors'])


def test_excessive_csv_rows_are_not_silently_truncated():
    content = '公司,年份,营业收入（亿元）\n' + '甲,2024,100\n' * 20001
    batch = prepare_import('过多行.csv', content.encode())
    assert not batch['can_commit']
    assert any('20,000 行' in error for error in batch['errors'])


def test_official_byd_pdf_financial_rows_and_year_columns():
    """Public 2024 report pp.10–11: previous-year delta is not a year column."""
    fixture = Path(__file__).parent / 'fixtures/byd_2024_key_financials.pdf'
    metadata = json.loads(fixture.with_suffix('.source.json').read_text(encoding='utf-8'))
    content = fixture.read_bytes()
    assert sha256(content).hexdigest() == metadata['fixture_sha256']
    batch = prepare_import(fixture.name, content, source_url=metadata['url'])
    assert batch['can_commit'], batch['errors']
    assert batch['companies'] == ['比亚迪股份有限公司']
    years = {record['year']: record for record in batch['records']}
    expected = {
        2024: (7771.02455, 402.54346, 1334.53873),
        2023: (6023.15354, 300.40811, 1697.25025),
        2022: (4240.60635, 166.22448, 1408.37657),
    }
    assert set(years) == set(expected)
    for year, values in expected.items():
        for metric, expected_value in zip(('revenue', 'net_profit', 'operating_cashflow'), values):
            assert years[year][metric] == pytest.approx(expected_value, abs=1e-8)
            source = years[year]['_provenance'][metric]
            assert source['raw_unit'] == '元'
            assert source['normalized_unit'] == '亿元'
            assert '%' not in source['raw_value']
            assert source['source_url'] == metadata['url']


def test_official_excerpt_provenance_keeps_local_pages_and_original_mapping():
    """The fixture is two pages; its pp.1/2 map explicitly to original pp.10/11."""
    fixture = Path(__file__).parent / 'fixtures/byd_2024_key_financials.pdf'
    metadata = json.loads(fixture.with_suffix('.source.json').read_text(encoding='utf-8'))
    _, records, warnings, _ = parse_pdf(fixture.name, fixture.read_bytes())
    current = next(record for record in records if record['year'] == 2024)
    assert not current['_errors'], warnings
    revenue_source = current['_provenance']['revenue']
    profit_source = current['_provenance']['net_profit']
    cashflow_source = current['_provenance']['operating_cashflow']
    assert revenue_source['page'] == 1
    assert profit_source['page'] == cashflow_source['page'] == 2
    assert metadata['page_offset'] == 9
    assert [revenue_source['page'] + metadata['page_offset'], profit_source['page'] + metadata['page_offset']] == metadata['original_pages'] == [10, 11]
    assert profit_source['cell'] != cashflow_source['cell']


def test_comparison_year_columns_do_not_register_other_year_reports(isolated):
    from core.retrieval import retrieve_documents
    fixture = Path(__file__).parent / 'fixtures/byd_2024_key_financials.pdf'
    batch = prepare_import(fixture.name, fixture.read_bytes())
    assert batch['report_year'] == 2024
    assert commit_import(batch, confirmed=True)['ok']
    reports = [dict(row) for row in db.fetch_report_files()]
    assert len(reports) == 1
    assert reports[0]['report_year'] == 2024
    assert '比较年度列不作为其他年度年报登记' in reports[0]['note']
    company_id = db.fetch_companies()[0]['id']
    rows = {row['year']: row for row in db.fetch_company_metrics(company_id)}
    assert rows[2023]['revenue'] == pytest.approx(6023.15354)
    assert rows[2022]['net_profit'] == pytest.approx(166.22448)
    unavailable = retrieve_documents('比亚迪2023年年报如何解释净利润变化？',
        {'companies': ['比亚迪股份有限公司'], 'years': [2023], 'metrics': ['net_profit']}, reports, isolated / 'rag')
    assert unavailable['status'] == 'no_evidence'
    assert unavailable['candidate_documents'] == 0
    assert not unavailable['citations']
    assert unavailable['missing_scope'] == [{'company': '比亚迪股份有限公司', 'year': 2023}]


def test_report_document_year_is_not_taken_from_hint_or_file_name(monkeypatch):
    import core.parsers as parsers
    # A financial table with a year column is not itself a titled annual report.
    text = '测试股份有限公司\n单位：亿元\n指标 2024年\n营业收入 100\n归母净利润 10\n经营现金流 20'
    monkeypatch.setattr(parsers, 'extract_pdf_text', lambda _: text)
    monkeypatch.setattr(parsers, '_pdf_pages', lambda _content, _text: [{'page': 1, 'text': text, 'tables': []}])
    batch = prepare_import('2024年年度报告.pdf', b'%PDF-placeholder', year_hint=2024)
    assert batch['can_commit'], batch['errors']
    assert batch['report_year'] is None
    assert batch['records'][0]['year'] == 2024
    assert any('文档不会冒充' in warning for warning in batch['warnings'])


def test_report_year_changes_invalidate_review_fingerprint(isolated):
    batch = prepare_import('2024年年度报告.pdf', make_pdf())
    assert batch['report_year'] == 2024
    batch['report_year'] = 2023
    assert not commit_import(batch, confirmed=True)['ok']
    assert not db.fetch_report_files()


def test_known_exchange_abbreviation_validates_without_renaming_pdf_company():
    fixture = Path(__file__).parent / 'fixtures/byd_2024_key_financials.pdf'
    batch = prepare_import(fixture.name, fixture.read_bytes(), expected_company='比亚迪', stock_code='002594',
                           report_title='比亚迪：2024年年度报告', year_hint=2024)
    assert batch['can_commit'], batch['errors']
    assert batch['companies'] == ['比亚迪股份有限公司']
    assert batch['stock_code'] == '002594'
    assert batch['report_year'] == 2024


@pytest.mark.parametrize(('expected', 'code'), [('贵州茅台', '600519'), (None, '600519'), ('比亚迪', '600519')])
def test_exchange_wrong_company_or_code_never_attaches_to_pdf(isolated, expected, code):
    fixture = Path(__file__).parent / 'fixtures/byd_2024_key_financials.pdf'
    batch = prepare_import(fixture.name, fixture.read_bytes(), expected_company=expected, stock_code=code)
    assert not batch['can_commit']
    assert any('股票代码' in error for error in batch['errors'])
    assert not commit_import(batch, confirmed=True)['ok']
    assert not db.fetch_companies()


def test_unknown_abbreviation_warns_but_does_not_replace_body_name():
    batch = prepare_import('测试2024年年度报告.pdf', make_pdf(), expected_company='测试')
    assert batch['can_commit'], batch['errors']
    assert batch['companies'] == ['测试股份有限公司']
    assert any('请人工核对' in warning for warning in batch['warnings'])
    unverified_code = prepare_import('测试2024年年度报告.pdf', make_pdf(), expected_company='测试', stock_code='600999')
    assert not unverified_code['can_commit']
    assert any('本地上传入口核验' in error for error in unverified_code['errors'])
    assert prepare_import('测试2024年年度报告.pdf', make_pdf())['can_commit']


def test_exchange_report_title_must_match_real_document_year():
    fixture = Path(__file__).parent / 'fixtures/byd_2024_key_financials.pdf'
    batch = prepare_import(fixture.name, fixture.read_bytes(), report_title='比亚迪2023年年度报告')
    assert not batch['can_commit']
    assert any('公告标题年度' in error for error in batch['errors'])


def test_true_xls_workbook_imports_multiple_sheets_with_units():
    fixture = Path(__file__).parent / 'fixtures/financial_import_multicompany.xls'
    # This fixture is a real BIFF8/OLE workbook, generated once using xlwt.
    assert fixture.read_bytes().startswith(bytes.fromhex('D0CF11E0A1B11AE1'))
    batch = prepare_import(fixture.name, fixture.read_bytes())
    assert batch['can_commit'], batch['errors']
    assert batch['companies'] == ['XLS乙企业', 'XLS甲企业']
    companies = {row['company_name']: row for row in batch['records']}
    assert companies['XLS甲企业']['revenue'] == 1
    assert companies['XLS乙企业']['revenue'] == 2
    assert companies['XLS甲企业']['net_profit'] == .1
    assert companies['XLS甲企业']['roe'] == 12.5  # underlying .125 in a real Excel percent cell
    assert companies['XLS乙企业']['roe'] == 12.5
